import json
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from gps_demo.onway import prepare, run_onway
from gps_demo.pipeline import ROOT


class OnwayTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / 'input').mkdir()
        self.config = json.loads((ROOT / 'config/onway.json').read_text(encoding='utf-8'))
        self.config.update(input='input', vehicles='master.xlsx', output='reports', routing='routing.json')
        self.config_path = self.root / 'config.json'
        self.config_path.write_text(json.dumps(self.config), encoding='utf-8')
        pd.DataFrame({'PLACA': ['AAA-123', 'BBB-456'], 'EMPRESA': ['Empresa A', 'Empresa B'], 'RUC': ['20111111111', '20222222222']}).to_excel(self.root / 'master.xlsx', index=False)
        self.frame = pd.DataFrame([
            [1, 'AAA-123', '15/10/2026', '08:00:00', 'SOS', '', '=malicioso'],
            [2, 'AAA-123', '15/10/2026', '09:00:00', 'Vehículo encendido', '', ''],
            [3, 'BBB-456', '15/10/2026', '10:00:00', 'Bache', '', ''],
            [4, 'AAA-123', '15/10/2026', '08:00:00', 'SOS', '', '=malicioso'],
            [5, 'BBB-456', '16/10/2026', '08:00:00', 'SOS', '', ''],
        ], columns=['Secuencia', 'Placa/Patente', 'Fecha', 'Hora', 'Alerta', 'Conductor', 'Notas'])
        self.source = self.root / 'input/one.xlsx'
        self.save()
        self.routing = {'allowed_recipients': ['51900000001', '51900000002'], 'coordinators': [
            {'id': 'one', 'alias': 'Prueba A', 'phone': '51900000001', 'company_rucs': ['20111111111']},
            {'id': 'two', 'alias': 'Prueba B', 'phone': '51900000002', 'company_rucs': ['20222222222']}]}
        self.save_routing()

    def save(self):
        self.frame.to_excel(self.source, index=False)
        # Guardar la muestra de fórmula como texto literal, no como fórmula
        # sin caché: así la prueba comprueba neutralización, no pérdida de datos.
        book = load_workbook(self.source)
        sheet = book.active
        for row_index, row in enumerate(self.frame.itertuples(index=False, name=None), 2):
            for col_index, value in enumerate(row, 1):
                if isinstance(value, str) and value.startswith('='):
                    sheet.cell(row_index, col_index).data_type = 's'
        book.save(self.source)
        book.close()

    def save_routing(self):
        (self.root / 'routing.json').write_text(json.dumps(self.routing))

    def test_filter_dedup_date_and_company(self):
        events, exported, stats = prepare(self.config, '2026-10-15', self.root)
        self.assertEqual((len(events), len(exported), stats['duplicados_exactos'], stats['eventos_operativos_excluidos']), (2, 3, 1, 1))
        self.assertEqual(set(events.Empresa), {'Empresa A', 'Empresa B'})
        self.assertEqual(set(events.Fecha), {pd.Timestamp('2026-10-15')})

    def test_unknown_alert_or_vehicle_blocks(self):
        for col, value in [('Alerta', 'Nueva alerta'), ('Placa/Patente', 'ZZZ-999'), ('Fecha', '2026-10-15')]:
            with self.subTest(col=col):
                original = self.frame.copy()
                self.frame.loc[0, col] = value
                self.save()
                with self.assertRaises(ValueError):
                    prepare(self.config, '2026-10-15', self.root)
                self.frame = original

    def test_reports_and_recipient_isolation(self):
        original_bytes = self.source.read_bytes()
        folder, manifest = run_onway(self.config_path, root=self.root)
        self.assertEqual(original_bytes, self.source.read_bytes())
        self.assertEqual(len(list(folder.glob('*.xlsx'))), 2)
        self.assertEqual(manifest['mensajes'][0]['documento'], 'EMPRESA_20111111111.xlsx')
        self.assertEqual(manifest['mensajes'][1]['documento'], 'EMPRESA_20222222222.xlsx')
        self.assertEqual(manifest['tipo_reporte'], 'empresa')
        self.assertFalse(list(folder.glob('COORDINADOR_*.xlsx')))
        self.assertNotIn('51900000001', json.dumps(manifest))
        for name, expected in [('EMPRESA_20111111111.xlsx', 'AAA-123'), ('EMPRESA_20222222222.xlsx', 'BBB-456')]:
            book = load_workbook(folder / name)
            self.assertEqual(book.sheetnames, ['Resumen', 'Report Alerts'])
            self.assertEqual(book['Resumen']['A6'].value, 1)
            self.assertEqual(len(book['Resumen']._charts), 1)
            sheet = book['Report Alerts']
            headers = [c.value for c in sheet[1]]
            self.assertNotIn('Conductor', headers)
            self.assertNotIn('Secuencia', headers)
            self.assertEqual(sheet.cell(2, headers.index('Placa/Patente') + 1).value, expected)
            self.assertTrue(all(c.data_type != 'f' for row in sheet for c in row))
            if expected == 'AAA-123':
                self.assertEqual(sheet.cell(2, headers.index('Notas') + 1).value, "'=malicioso")
            self.assertIsInstance(sheet.cell(2, headers.index('Fecha') + 1).value, datetime)
            book.close()

    def test_no_selected_alerts_has_explicit_zero(self):
        self.frame = self.frame.iloc[[1]].copy()
        self.save()
        folder, manifest = run_onway(self.config_path, root=self.root)
        self.assertEqual(manifest['alertas_reportadas'], 0)
        book = load_workbook(folder / 'EMPRESA_20111111111.xlsx')
        self.assertEqual(book['Resumen']['A6'].value, 0)
        self.assertIn('Sin alertas', book['Resumen']['A13'].value)
        self.assertEqual(book['Report Alerts'].max_row, 1)
        book.close()

    def test_routing_validation_prevents_generation(self):
        self.routing['coordinators'][0]['phone'] = '51911111111'
        self.save_routing()
        with self.assertRaises(ValueError):
            run_onway(self.config_path, root=self.root)
        self.assertFalse((self.root / 'reports').exists())

    def test_duplicate_company_routing_and_live_mode(self):
        self.routing['coordinators'][1]['company_rucs'] = ['20111111111']
        self.save_routing()
        with self.assertRaises(ValueError):
            run_onway(self.config_path, root=self.root)
        with patch.dict(os.environ, {'WHATSAPP_MODE': 'live'}):
            with self.assertRaises(ValueError):
                run_onway(self.config_path, root=self.root)

    def test_trial_rejects_multiple_companies_or_same_number(self):
        self.routing['coordinators'][0]['company_rucs'] = ['20111111111', '20333333333']
        self.save_routing()
        with self.assertRaises(ValueError):
            run_onway(self.config_path, root=self.root)
        self.routing['coordinators'][0]['company_rucs'] = ['20111111111']
        self.routing['coordinators'][1]['phone'] = self.routing['coordinators'][0]['phone']
        self.save_routing()
        with self.assertRaises(ValueError):
            run_onway(self.config_path, root=self.root)


if __name__ == '__main__':
    unittest.main()
