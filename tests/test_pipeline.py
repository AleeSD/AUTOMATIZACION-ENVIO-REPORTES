import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from gps_demo.pipeline import ROOT, prepare, run


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.columns = {x: x for x in ['placa', 'fecha', 'alerta', 'conductor']}
        self.source = ROOT / 'data/samples/alertas.csv'
        self.assignments = ROOT / 'data/samples/asignaciones.csv'

    def parse(self, source=None, mapping=None, day='2026-10-16'):
        return prepare(source or self.source, mapping or self.assignments, self.columns, day)

    def test_counts_and_deduplication(self):
        events, duplicates = self.parse()
        self.assertEqual((len(events), duplicates), (4, 1))
        self.assertEqual(events.groupby('coordinador_id').size().to_dict(), {'coord_norte': 3, 'coord_sur': 1})

    def test_xlsx_and_semicolon_equivalence(self):
        source = pd.read_csv(self.source)
        for suffix in ['xlsx', 'csv']:
            path = self.folder / ('input.' + suffix)
            if suffix == 'xlsx':
                source.to_excel(path, index=False, engine='openpyxl')
            else:
                source.to_csv(path, index=False, sep=';')
            pd.testing.assert_frame_equal(self.parse()[0], self.parse(path)[0])

    def test_invalid_rows_block(self):
        for column, value in [('placa', 'ZZZ999'), ('fecha', '16/10/2026'), ('alerta', '')]:
            with self.subTest(column=column):
                frame = pd.read_csv(self.source)
                frame.loc[0, column] = value
                path = self.folder / 'bad.csv'
                frame.to_csv(path, index=False)
                with self.assertRaises(ValueError):
                    self.parse(path)

    def test_duplicate_assignment_and_empty_day(self):
        frame = pd.read_csv(self.assignments, dtype=str)
        path = self.folder / 'mapping.csv'
        pd.concat([frame, frame.iloc[:1]]).to_csv(path, index=False)
        with self.assertRaises(ValueError):
            self.parse(mapping=path)
        with self.assertRaises(ValueError):
            self.parse(day='2026-10-17')

    def test_reports_isolation_and_formula_safety(self):
        frame = pd.read_csv(self.source)
        frame.loc[0, 'conductor'] = '=HYPERLINK("https://example.com")'
        source = self.folder / 'source.csv'
        frame.to_csv(source, index=False)
        config = json.loads((ROOT / 'config/demo.json').read_text())
        config['output'] = str(self.folder / 'reports')
        config_path = self.folder / 'config.json'
        config_path.write_text(json.dumps(config))
        folder, manifest = run(config_path, source)
        self.assertEqual(len(list(folder.glob('*.xlsx'))), 5)
        self.assertEqual(manifest['modo'], 'simulation')
        self.assertNotIn('telefono', json.dumps(manifest))
        for name, plates, count in [('coord_norte', {'ABC123', 'DEF456'}, 3), ('coord_sur', {'GHI789'}, 1)]:
            book = load_workbook(folder / f'coordinador_{name}.xlsx')
            rows = list(book['Detalle'].iter_rows(min_row=2, values_only=True))
            self.assertEqual({r[0] for r in rows}, plates)
            self.assertEqual(len(rows), count)
            self.assertEqual(book['Resumen']['B3'].value, count)
            self.assertTrue(all(c.data_type != 'f' for row in book['Detalle'] for c in row))
            book.close()

    def test_real_mode_blocked(self):
        with patch.dict(os.environ, {'WHATSAPP_MODE': 'live'}):
            with self.assertRaises(ValueError):
                run()


if __name__ == '__main__':
    unittest.main()
