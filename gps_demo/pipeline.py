import argparse
import hmac
import json
import os
import re
import uuid
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]


def plate(value):
    value = re.sub(r'[\s-]', '', str(value)).upper()
    if not re.fullmatch(r'[A-Z0-9]{5,10}', value):
        raise ValueError('Placa inválida')
    return value


def read_table(path):
    path = Path(path)
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('Archivo supera 10 MB')
    if path.suffix.lower() == '.csv':
        frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding='utf-8-sig', sep=None, engine='python')
    elif path.suffix.lower() == '.xlsx':
        frame = pd.read_excel(path, dtype=str, keep_default_na=False, engine='openpyxl')
    else:
        raise ValueError('Solo CSV o XLSX')
    if frame.empty or len(frame) > 50000:
        raise ValueError('Archivo vacío o supera 50 000 filas')
    return frame


def require(frame, columns):
    if not set(columns).issubset(frame.columns):
        raise ValueError('Faltan columnas: ' + ', '.join(sorted(set(columns) - set(frame.columns))))
    frame = frame[list(columns)].copy()
    for col in columns:
        frame[col] = frame[col].astype(str).str.strip()
        if frame[col].eq('').any():
            raise ValueError('Campo vacío: ' + col)
    return frame


def prepare(input_path, assignments_path, columns, day):
    events = read_table(input_path)
    events = require(events, list(columns.values())).rename(columns={v: k for k, v in columns.items()})
    events['placa'] = events['placa'].map(plate)
    if not events.fecha.str.fullmatch(r'\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}').all():
        raise ValueError('Fecha debe ser ISO AAAA-MM-DD HH:MM:SS, hora Lima')
    events['fecha'] = pd.to_datetime(events.fecha, format='%Y-%m-%d %H:%M:%S', errors='raise')
    original = len(events)
    events = events.drop_duplicates()
    duplicates = original - len(events)
    mapping = require(read_table(assignments_path), ['placa', 'coordinador_id', 'coordinador', 'telefono'])
    mapping['placa'] = mapping.placa.map(plate)
    if mapping.placa.duplicated().any():
        raise ValueError('Placa duplicada en asignaciones')
    if not mapping.coordinador_id.str.fullmatch(r'[a-zA-Z0-9_-]{1,60}').all():
        raise ValueError('coordinador_id inválido')
    if not mapping.telefono.str.fullmatch(r'\d{8,15}').all():
        raise ValueError('Teléfono inválido')
    if (mapping.groupby('coordinador_id')[['coordinador', 'telefono']].nunique() > 1).any().any():
        raise ValueError('Datos contradictorios de coordinador')
    unknown = set(events.placa) - set(mapping.placa)
    if unknown:
        raise ValueError('Placas sin asignación: ' + ', '.join(sorted(unknown)))
    events = events[events.fecha.dt.date == date.fromisoformat(day)]
    if events.empty:
        raise ValueError('Sin eventos para el día solicitado')
    return events.merge(mapping, on='placa', validate='many_to_one').sort_values(['fecha', 'placa']), duplicates


def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def workbook(path, frame, title, day):
    book = Workbook()
    summary = book.active
    summary.title = 'Resumen'
    summary.append(['O.S. Logistics Perú S.A.C.', title])
    summary.append(['Fecha (Lima)', day])
    summary.append(['Eventos registrados', len(frame)])
    summary.append(['Fuente', 'Exportación manual; muestra si se usa configuración demo'])
    summary.append(['Nota', 'Conteos de alertas; no es evaluación del conductor'])
    summary.append(['Placa', 'Alerta', 'Cantidad'])
    for (p, alert), count in frame.groupby(['placa', 'alerta']).size().items():
        summary.append([safe(p), safe(alert), int(count)])
    detail = book.create_sheet('Detalle')
    cols = ['placa', 'fecha', 'alerta', 'conductor', 'coordinador']
    detail.append(cols)
    for row in frame[cols].itertuples(index=False, name=None):
        detail.append([safe(v) for v in row])
    for sheet, header in [(summary, 6), (detail, 1)]:
        sheet.freeze_panes = f'A{header + 1}'
        sheet.auto_filter.ref = f'A{header}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}'
        for cell in sheet[header]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='163B59')
        for column in sheet.columns:
            width = min(65, max(16, max(len(str(c.value or '')) for c in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
            for cell in column:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = 'dd/mm/yyyy hh:mm:ss'
        for row in range(1, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 32
    book.save(path)


def run(config_path=ROOT / 'config/demo.json', input_path=None, day=None):
    if os.environ.get('WHATSAPP_MODE', 'simulation') != 'simulation':
        raise ValueError('Envío real bloqueado en esta versión')
    config = json.loads(Path(config_path).read_text(encoding='utf-8'))
    if config.get('parser') == 'onway':
        from .onway import run_onway
        return run_onway(config_path, day)
    selected_day = day or config['date']
    events, duplicates = prepare(input_path or ROOT / config['input'], ROOT / config['assignments'], config['columns'], selected_day)
    folder = ROOT / config['output'] / (datetime.now().strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:8])
    folder.mkdir(parents=True, exist_ok=False)
    for p, frame in events.groupby('placa'):
        workbook(folder / f'placa_{p}.xlsx', frame, 'Placa ' + p, selected_day)
    messages = []
    for coordinator, frame in events.groupby('coordinador_id'):
        filename = f'coordinador_{coordinator}.xlsx'
        workbook(folder / filename, frame, 'Consolidado ' + coordinator, selected_day)
        messages.append({'coordinador_id': coordinator, 'documento': filename, 'estado': 'simulado',
                         'texto': f'{selected_day}: {len(frame)} alertas registradas en {frame.placa.nunique()} placas. Revisar el detalle adjunto; el conteo no califica al conductor.'})
    manifest = {'modo': 'simulation', 'fecha': selected_day, 'eventos': len(events), 'duplicados_eliminados': duplicates, 'mensajes': messages}
    (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    return folder, manifest


def serve(config_path=ROOT / 'config/onway.json'):
    token = os.environ.get('GPS_LOCAL_TOKEN', '')
    if len(token) < 24:
        raise ValueError('GPS_LOCAL_TOKEN debe tener al menos 24 caracteres')
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            if self.path != '/run':
                self.send_error(404)
                return
            if not hmac.compare_digest(self.headers.get('X-GPS-Token', ''), token):
                self.send_error(403)
                return
            try:
                folder, manifest = run(config_path)
                payload = json.dumps({'carpeta': str(folder), **manifest}, ensure_ascii=False).encode()
                self.send_response(200)
            except Exception:
                payload = b'{"error":"Procesamiento fallido; revisar entrada y configuracion local"}'
                self.send_response(422)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
        def log_message(self, *args):
            pass
    print('Receptor local: http://127.0.0.1:8765 ; Ctrl+C para detener')
    HTTPServer(('127.0.0.1', 8765), Handler).serve_forever()


def main():
    parser = argparse.ArgumentParser(description='Demo local GPS; simulación obligatoria')
    parser.add_argument('command', choices=['run', 'serve', 'onway'])
    parser.add_argument('--input', type=Path)
    parser.add_argument('--date')
    parser.add_argument('--config', type=Path, default=ROOT / 'config/demo.json')
    args = parser.parse_args()
    try:
        if args.command == 'serve':
            serve(args.config if args.config != ROOT / 'config/demo.json' else ROOT / 'config/onway.json')
        elif args.command == 'onway':
            if not args.date:
                raise ValueError('Onway requiere --date AAAA-MM-DD')
            from .onway import run_onway
            folder, manifest = run_onway(day=args.date)
            print(json.dumps({'carpeta': str(folder), 'alertas': manifest['alertas_reportadas'], 'reportes_empresa': len(manifest['documentos_empresa']), 'mensajes_simulados': len(manifest['mensajes'])}, ensure_ascii=False, indent=2))
        else:
            if args.input and not args.date:
                raise ValueError('Archivo propio requiere --date AAAA-MM-DD')
            folder, manifest = run(args.config, args.input, args.date)
            print(json.dumps({'carpeta': str(folder), **manifest}, ensure_ascii=False, indent=2))
    except (ValueError, OSError, KeyError) as error:
        parser.exit(1, f'Error: {error}\n')
