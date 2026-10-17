"""Importación Onway y reportes locales. Sin cliente de envío externo."""
import hashlib
import json
import os
import re
import uuid
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import ROOT, plate, safe

NAVY = '163B59'
TEAL = '087F8C'


def read(path):
    path = Path(path)
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError(f'Archivo supera 10 MB: {path.name}')
    if path.suffix.lower() == '.csv':
        frame = pd.read_csv(path, sep=None, engine='python', encoding='utf-8-sig', keep_default_na=False)
    elif path.suffix.lower() == '.xlsx':
        frame = pd.read_excel(path, engine='openpyxl', keep_default_na=False)
    else:
        raise ValueError('Solo CSV o XLSX')
    if len(frame) > 50000 or frame.empty:
        raise ValueError(f'Archivo vacío o supera 50 000 filas: {path.name}')
    if frame.columns.duplicated().any():
        raise ValueError('Columnas duplicadas')
    for c in frame.select_dtypes(include='object'):
        frame[c] = frame[c].map(lambda v: v.strip() if isinstance(v, str) else v)
    return frame


def mandatory(frame, names):
    if not set(names).issubset(frame.columns):
        raise ValueError('Faltan columnas: ' + ', '.join(set(names) - set(frame.columns)))
    for name in names:
        if frame[name].isna().any() or frame[name].astype(str).str.strip().eq('').any():
            raise ValueError('Campo vacío: ' + name)


def onway_date(value):
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value), '%d/%m/%Y').date()


def onway_time(value):
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.time()
    if isinstance(value, time):
        return value
    return datetime.strptime(str(value), '%H:%M:%S').time()


def prepare(config, day, root=ROOT):
    selected = date.fromisoformat(day)
    master = read(root / config['vehicles'])
    mandatory(master, ['PLACA', 'EMPRESA', 'RUC'])
    master['_placa'] = master.PLACA.map(plate)
    master['RUC'] = master.RUC.astype(str)
    if master._placa.duplicated().any() or not master.RUC.str.fullmatch(r'\d{11}').all():
        raise ValueError('Padrón: placa duplicada o RUC inválido')
    if (master.groupby('RUC').EMPRESA.nunique() > 1).any():
        raise ValueError('Padrón: empresas contradictorias para un RUC')
    source = root / config['input']
    paths = sorted(p for p in source.rglob('*') if p.suffix.lower() in {'.xlsx', '.csv'} and not p.name.startswith('~$')) if source.is_dir() else [source]
    if not paths:
        raise ValueError('No hay reportes Onway')
    frames, audit = [], []
    for path in paths:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        frame = read(path)
        mandatory(frame, ['Placa/Patente', 'Fecha', 'Hora', 'Alerta'])
        frame['_placa'] = frame['Placa/Patente'].map(plate)
        frame['_dia'] = frame.Fecha.map(onway_date)
        frame['Hora'] = frame.Hora.map(onway_time)
        frame = frame[frame._dia == selected].copy()
        audit.append({'archivo': str(path.relative_to(root)), 'sha256': digest, 'filas_dia': len(frame)})
        if not frame.empty:
            frames.append(frame)
    if not frames:
        raise ValueError('Sin eventos para el día seleccionado')
    all_events = pd.concat(frames, ignore_index=True).fillna('')
    if len(all_events) > 50000:
        raise ValueError('Lote supera 50 000 filas')
    unknown_plates = set(all_events._placa) - set(master._placa)
    if unknown_plates:
        raise ValueError('Placas sin padrón: ' + ', '.join(sorted(unknown_plates)))
    included = set(config['included_alerts'])
    excluded = set(config['excluded_alerts'])
    if included & excluded:
        raise ValueError('Reglas contradictorias de alertas')
    unknown_alerts = set(all_events.Alerta) - included - excluded
    if unknown_alerts:
        raise ValueError('Clasificar nuevas alertas en config/onway.json: ' + ', '.join(sorted(unknown_alerts)))
    rows_before = len(all_events)
    # No deduplicar por campos eliminados: dos eventos con fecha de atención
    # diferente pueden ser legítimos. Solo omitir secuencia y diferencias de placa.
    keys = [c for c in all_events if c not in {'Secuencia', 'Placa/Patente'}]
    all_events = all_events.drop_duplicates(subset=keys)
    duplicates = rows_before - len(all_events)
    lookup = master.set_index('_placa')
    all_events['Placa/Patente'] = all_events._placa.map(lookup.PLACA)
    all_events['Empresa'] = all_events._placa.map(lookup.EMPRESA)
    all_events['_ruc'] = all_events._placa.map(lookup.RUC)
    all_events['Fecha'] = all_events._dia.map(lambda d: datetime.combine(d, time()))
    events = all_events[all_events.Alerta.isin(included)].sort_values(['_placa', 'Fecha', 'Hora']).copy()
    stats = {'filas_originales': rows_before, 'duplicados_exactos': duplicates,
             'eventos_operativos_excluidos': len(all_events) - len(events), 'alertas_reportadas': len(events),
             'placas_con_exportacion': sorted(all_events['Placa/Patente'].unique().tolist()),
             'clasificacion': {'incluidas': sorted(included), 'excluidas': sorted(excluded)}, 'fuentes': audit}
    return events, all_events, stats


def styling(sheet, header=1):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f'A{header + 1}'
    for cell in sheet[header]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    sheet.row_dimensions[header].height = 32
    for row in sheet.iter_rows(min_row=header + 1):
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color=NAVY)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EFF5F8')
        sheet.row_dimensions[row[0].row].height = 32


def write_report(path, events, exported, config, title, day):
    book = Workbook()
    dash = book.active
    dash.title = 'Resumen'
    dash.sheet_view.showGridLines = False
    dash.merge_cells('A1:L2')
    dash['A1'] = safe('O.S. LOGISTICS PERÚ | ' + title)
    dash['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    dash['A1'].fill = PatternFill('solid', fgColor=NAVY)
    dash['A1'].alignment = Alignment(vertical='center', wrap_text=True)
    dash.row_dimensions[1].height = 28
    dash.row_dimensions[2].height = 28
    dash.merge_cells('A3:L3')
    dash['A3'] = f'{day} · Solo conducción y seguridad · Hora local Lima'
    dash['A3'].font = Font(name='Calibri', size=12, color=TEAL)
    speed = int(events.Alerta.str.contains('velocidad', case=False).sum())
    harsh = int(events.Alerta.str.contains('brusc', case=False).sum())
    sos = int(events.Alerta.eq('SOS').sum())
    for first, last, label, value in [('A', 'C', 'ALERTAS', len(events)), ('D', 'F', 'VELOCIDAD', speed), ('G', 'I', 'MANIOBRAS BRUSCAS', harsh), ('J', 'L', 'SOS', sos)]:
        dash.merge_cells(f'{first}5:{last}5')
        dash.merge_cells(f'{first}6:{last}7')
        dash[f'{first}5'] = label
        dash[f'{first}5'].font = Font(name='Calibri', size=11, bold=True, color=TEAL)
        cell = dash[f'{first}6']
        cell.value = value
        cell.font = Font(name='Calibri', size=28, bold=True, color=NAVY if label != 'SOS' or value == 0 else 'B42318')
        cell.fill = PatternFill('solid', fgColor='E8F3F5')
    dash.merge_cells('A9:L9')
    dash['A9'] = f'{exported._placa.nunique()} placas con exportación · {events._placa.nunique()} con alertas seleccionadas. Sin puntaje del conductor.'
    dash.merge_cells('A10:L10')
    dash['A10'] = 'Principales alertas por frecuencia; SOS se destaca por separado. Revisar el detalle antes de tomar acciones.'
    dash['A10'].alignment = Alignment(wrap_text=True)
    dash.row_dimensions[10].height = 32
    for col, label in [('A', 'TIPO DE ALERTA'), ('E', 'CANTIDAD')]:
        dash[f'{col}12'] = label
        dash[f'{col}12'].font = Font(bold=True, color='FFFFFF')
        dash[f'{col}12'].fill = PatternFill('solid', fgColor=NAVY)
    dash.merge_cells('A12:D12')
    counts = events.Alerta.value_counts().sort_values(ascending=False)
    for row, (alert, count) in enumerate(counts.items(), 13):
        dash.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        dash.cell(row, 1, safe(alert)).alignment = Alignment(wrap_text=True, vertical='center')
        dash.cell(row, 5, int(count))
        dash.row_dimensions[row].height = 32
    if not counts.empty:
        chart = BarChart()
        chart.type = 'bar'
        chart.style = 10
        chart.title = 'Alertas registradas'
        chart.add_data(Reference(dash, min_col=5, min_row=12, max_row=12 + len(counts)), titles_from_data=True)
        chart.set_categories(Reference(dash, min_col=1, min_row=13, max_row=12 + len(counts)))
        chart.legend = None
        chart.height = 9
        chart.width = 20
        chart.y_axis.title = 'Cantidad'
        dash.add_chart(chart, 'G12')
    else:
        dash.merge_cells('A13:E14')
        dash['A13'] = 'Sin alertas de las categorías seleccionadas en la exportación. No acredita ausencia de incidentes.'
        dash['A13'].alignment = Alignment(wrap_text=True)
    start = max(32, 15 + len(counts))
    dash.cell(start, 1, 'PLACA')
    dash.cell(start, 3, 'EMPRESA')
    dash.cell(start, 10, 'ALERTAS')
    for row, (p, data) in enumerate(exported.groupby('_placa'), start + 1):
        dash.cell(row, 1, safe(str(data['Placa/Patente'].iloc[0])))
        dash.merge_cells(start_row=row, start_column=3, end_row=row, end_column=9)
        dash.cell(row, 3, safe(str(data.Empresa.iloc[0]))).alignment = Alignment(wrap_text=True)
        dash.cell(row, 10, int(events._placa.eq(p).sum()))
        dash.row_dimensions[row].height = 32
    dash.merge_cells(start_row=start + exported._placa.nunique() + 3, start_column=1,
                     end_row=start + exported._placa.nunique() + 4, end_column=12)
    note = dash.cell(start + exported._placa.nunique() + 3, 1, 'Fuente: exportación manual Onway y UNIDADES Y RUC.xlsx. Se excluyen eventos operativos. Clasificación configurable; no es evaluación de responsabilidad.')
    note.alignment = Alignment(wrap_text=True, vertical='center')
    for i in range(1, 13):
        dash.column_dimensions[get_column_letter(i)].width = 12
    dash.freeze_panes = 'A12'
    dash.sheet_properties.pageSetUpPr.fitToPage = True
    dash.page_setup.orientation = 'landscape'
    dash.page_setup.paperSize = dash.PAPERSIZE_A3
    dash.page_setup.fitToWidth = 1
    dash.page_setup.fitToHeight = 1
    dash.print_options.horizontalCentered = True
    dash.print_area = f'A1:Q{max(36, start + exported._placa.nunique() + 5)}'
    detail = book.create_sheet('Report Alerts')
    # Orden original de Onway, sin columnas administrativas ni vacías.
    cols = [c for c in exported if c in config['detail_columns'] and (c in {'Placa/Patente', 'Fecha', 'Hora', 'Alerta', 'Empresa'} or (not events.empty and events[c].astype(str).str.strip().ne('').any()))]
    detail.append(cols)
    for record in events[cols].itertuples(index=False, name=None):
        detail.append([safe(v) for v in record])
    styling(detail)
    detail.auto_filter.ref = detail.dimensions
    detail.print_title_rows = '1:1'
    detail.sheet_properties.pageSetUpPr.fitToPage = True
    detail.page_setup.orientation = 'landscape'
    detail.page_setup.fitToWidth = 1
    detail.page_setup.fitToHeight = 0
    for i, name in enumerate(cols, 1):
        detail.column_dimensions[get_column_letter(i)].width = 48 if name == 'Empresa' else 36 if name in {'Alerta', 'Alerta alias', 'Dirección', 'Notas'} else max(18, min(30, len(name) + 2))
        for column_cell in list(detail.columns)[i - 1][1:]:
            if name == 'Fecha':
                column_cell.number_format = 'dd/mm/yyyy'
            elif name == 'Hora':
                column_cell.number_format = 'hh:mm:ss'
            elif name in {'Latitud', 'Longitud'}:
                column_cell.number_format = '0.000000'
            elif isinstance(column_cell.value, (int, float)):
                column_cell.number_format = '0.0'
    book.save(path)
    return cols


def routing(config, companies, root=ROOT):
    path = root / config['routing']
    if not path.exists():
        return [], sorted(companies)
    data = json.loads(path.read_text(encoding='utf-8'))
    allowlist = set(data['allowed_recipients'])
    if not all(re.fullmatch(r'51\d{9}', n) for n in allowlist):
        raise ValueError('Número de prueba inválido')
    assigned, entries, ids, phones = set(), [], set(), set()
    if len(data['coordinators']) > 3:
        raise ValueError('Demo limitada a tres coordinadores')
    for entry in data['coordinators']:
        if entry['id'] in ids or not re.fullmatch(r'[A-Za-z0-9_-]{1,50}', entry['id']):
            raise ValueError('Identificador de coordinador inválido o duplicado')
        ids.add(entry['id'])
        if entry['phone'] not in allowlist:
            raise ValueError('Destinatario fuera de lista autorizada')
        if entry['phone'] in phones:
            raise ValueError('La demo admite un único reporte por número')
        phones.add(entry['phone'])
        rucs = set(entry['company_rucs'])
        if len(rucs) != 1:
            raise ValueError('Seleccionar exactamente una empresa por coordinador de prueba')
        if rucs & assigned:
            raise ValueError('Empresa asignada a más de un coordinador de prueba')
        assigned |= rucs
        entries.append(entry)
    return entries, sorted(companies - assigned)


def run_onway(config_path=ROOT / 'config/onway.json', day=None, root=ROOT, source_dir=None):
    if os.environ.get('WHATSAPP_MODE', 'simulation') != 'simulation':
        raise ValueError('Envío real bloqueado; integración Meta aún no configurada')
    config = json.loads(Path(config_path).read_text(encoding='utf-8'))
    if source_dir is not None:
        config['input'] = str(source_dir)
    day = day or config['date']
    events, exported, stats = prepare(config, day, root)
    entries, pending = routing(config, set(exported._ruc), root)
    folder = root / config['output'] / day / uuid.uuid4().hex[:10]
    folder.mkdir(parents=True, exist_ok=False)
    documents = []
    for ruc, data in exported.groupby('_ruc'):
        selected = events[events._ruc == ruc]
        filename = f'EMPRESA_{ruc}.xlsx'
        cols = write_report(folder / filename, selected, data, config, str(data.Empresa.iloc[0]), day)
        document = {'archivo': filename, 'ruc': str(ruc), 'empresa': str(data.Empresa.iloc[0]),
                          'placas': sorted(data['Placa/Patente'].unique().tolist()), 'alertas': len(selected),
                          'columnas': cols, 'sha256': hashlib.sha256((folder / filename).read_bytes()).hexdigest()}
        if config.get('pdf_reports'):
            from .pdf_report import write_pdf
            pdf_name = f'EMPRESA_{ruc}.pdf'
            write_pdf(folder / pdf_name, selected, data, str(data.Empresa.iloc[0]), day)
            document['pdf'] = {'archivo': pdf_name, 'sha256': hashlib.sha256((folder / pdf_name).read_bytes()).hexdigest()}
        documents.append(document)
    messages = []
    for entry in entries:
        data = exported[exported._ruc.isin(entry['company_rucs'])]
        if data.empty:
            continue
        selected = events[events._ruc.isin(entry['company_rucs'])]
        ruc = entry['company_rucs'][0]
        filename = f'EMPRESA_{ruc}.xlsx'
        company = str(data.Empresa.iloc[0])
        messages.append({'coordinador_id': entry['id'], 'alias': entry['alias'], 'estado': 'simulado',
                         'telefono_mascarado': '***' + entry['phone'][-3:],
                         'empresa_ruc': ruc, 'empresa': company, 'documento': filename,
                         'texto': f'{company} | {day}: {len(selected)} alertas de conducción y seguridad en {data._placa.nunique()} placas con exportación. Velocidad: {int(selected.Alerta.str.contains("velocidad", case=False).sum())}; SOS: {int(selected.Alerta.eq("SOS").sum())}. Revisar el detalle; no es calificación del conductor.'})
    manifest = {'schema_version': 2, 'tipo_reporte': 'empresa', 'modo': 'simulation', 'fecha': day, **stats, 'documentos_empresa': documents,
                'empresas_no_seleccionadas_demo': pending, 'mensajes': messages}
    (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    reference = root / config['output'] / 'ultimo_lote_empresa.json'
    temporary = reference.with_name('.ultimo_lote_' + uuid.uuid4().hex + '.tmp')
    temporary.write_text(json.dumps({'fecha': day, 'lote': str(folder.relative_to(root / config['output']))}), encoding='utf-8')
    temporary.replace(reference)
    return folder, manifest
